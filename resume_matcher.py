import os
import json
import time
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, field_validator
from pypdf import PdfReader
from docx import Document

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")
if not my_api_key:
    raise ValueError("groq api key is not defined in .env file")

client = Groq(api_key=my_api_key)
model = "openai/gpt-oss-20b"

job_description = """
# Job Title: Python Developer

**Company:** TechNova Solutions  
**Location:** Remote  
**Job Type:** Full-time  
**Experience Required:** 2+ years

### Job Summary
We are looking for a skilled Python Developer to develop, maintain, and optimize backend applications and REST APIs. The ideal candidate should have experience with Python, FastAPI, databases, and software development best practices.

### Required Skills
- Python programming
- FastAPI or Flask
- SQL
- REST API development
- Git and version control
- Problem-solving and debugging

### Preferred Skills
- Docker
- AWS cloud services
- PostgreSQL
- Pytest and unit testing
- Basic knowledge of CI/CD

### Responsibilities
- Develop and maintain backend applications using Python.
- Build and integrate REST APIs using FastAPI.
- Write efficient SQL queries and optimize database performance.
- Write unit tests and debug application issues.
- Collaborate with frontend developers and other team members.
- Use Git for version control and code collaboration.
- Follow coding standards and security best practices.

### Education Requirements
Bachelor's degree in Computer Science, Information Technology, Software Engineering, or a related field.

### Experience Requirements
At least 2 years of relevant Python development experience, including backend development and API integration.

### Certifications
Not mandatory. Relevant Python or cloud certifications are an advantage.

"""

# ---------- Models ----------
class JobD(BaseModel):
    role: str
    required_skill: list[str] = []
    preferred_skill: list[str] = []
    minimum_experiences: float | None = None
    education_requirements: list[str] = []
    responsibilities: list[str] = []

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    total_experience_years: float | None = None
    skill: list[str] = []
    experience: list[str] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []

    # agar model null bheje to empty list bana do
    @field_validator("skill", "experience", "education", "projects",
                     "certifications", mode="before")
    @classmethod
    def none_to_list(cls, v):
        return v or []

class MatchResult(BaseModel):
    candidate_name: str | None = None
    matching_skills: list[str] = []
    missing_skills: list[str] = []
    experience_requirement_met: bool | None = None
    score: float
    verdict: str = ""

# ---------- LLM helper ----------
def ask_json(system, user):
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
        response_format={"type": "json_object"},
    )
    return json.loads(response.choices[0].message.content)

# ---------- Job description parse ----------
jobd_schema = JobD.model_json_schema()
job_system = f"""
You are an AI assistant that extracts structured information from a job description.
Use only information present in the job description. Do not invent missing information.
Return JSON matching this schema:
{jobd_schema}
"""
job = JobD(**ask_json(job_system, f"Analyze the following job description:\n{job_description}"))

# ---------- Resume parse ----------
resume_schema = Resume.model_json_schema()

def parse_resume(resume_text):
    system = f"""
    You are an expert resume parser.
    Extract information based on meaning, not only exact section headings.
    Headings like experience, professional experience, work history,
    employment, internship all count as experience.
    Skills may appear in skills section, work experience, internship or projects.
    Return only valid JSON matching this schema:
    {resume_schema}
    Rules: do not invent information. Missing value -> null. Empty list -> [].
    Include internships under experience. Extract skills from the entire resume.
    """
    data = ask_json(system, f"Parse the following resume:\n{resume_text}")
    return Resume(**data)

# ---------- Scoring ----------
def final_score(job, resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are an HR recruiter.
    Compare the candidate resume with the job description.

    Job description:
    {job.model_dump_json(indent=2)}

    Candidate resume:
    {resume.model_dump_json(indent=2)}

    Return JSON matching this schema:
    {match_schema}

    Fill: candidate name, matching skills, missing important skills,
    whether experience requirement is met, overall match score (0 to 100),
    and a short final verdict. Keep it concise.
    """
    data = ask_json("You are an HR recruiter. Reply only in JSON.", prompt)
    return MatchResult(**data)

# ---------- File readers ----------
def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text

def read_resume(file_path):
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(file_path)
    elif suffix == ".docx":
        return read_docx(file_path)
    return None

# ---------- Main ----------
resume_folder = Path(__file__).parent / "resumes"
all_result = []

for file_path in resume_folder.iterdir():
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue
    print("\nProcessing", file_path.name)

    resume_text = read_resume(file_path)
    if not resume_text or not resume_text.strip():
        print("Text nahi mila (scanned PDF ho sakta hai), skip.")
        continue

    try:
        parsed = parse_resume(resume_text)
        time.sleep(5)
        result = final_score(job, parsed)
        time.sleep(5)
    except Exception as e:
        print("Error:", e)
        continue

    print("score:", result.score)
    all_result.append({
        "name": parsed.name or result.candidate_name,
        "score": result.score,
        "details": result.model_dump(),
    })

all_result.sort(key=lambda c: c["score"], reverse=True)

print("\nTOP 2 CANDIDATES")
for candidate in all_result[:2]:
    print(candidate["name"], "-", candidate["score"], "%")
    print(candidate["details"])